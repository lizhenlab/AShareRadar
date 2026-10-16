"""Saved rule exports keep the complete frozen rule rather than overlaying v6."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
import json

from openpyxl import load_workbook
import pytest

from app.market_scan_screening import screen_spec_from_discovery
from app.models.discovery import DiscoveryPresetCreate
from app.services.market_scan_export import MarketScanExportFilters, build_market_scan_workbook, frozen_market_scan_export_filters
from tests.test_market_scan_ranking_maintenance_snapshot import (
    _CapturedRankingStore, _query_service, published_ranking_source as _published_source,
)
from tests.test_market_scan_selection import _client, _members, _selection


@pytest.fixture(scope="module")
def ranking_source(tmp_path_factory):
    return _published_source.__wrapped__(tmp_path_factory)


def test_frozen_preset_export_keeps_base_membership_when_production_v6_changes_scores(ranking_source):
    cache, run, context, records = ranking_source
    _members(cache, list(records))
    cutoff = next(n for n in range(101) if {s for s, r in records.items() if r["base_score"] >= n} != {s for s, r in records.items() if r["score"] >= n})
    service = cache.domain_services.discovery
    preset = service.create_preset(DiscoveryPresetCreate(name="v6隔离", criteria={"score": {"min": cutoff}}))
    displayed = service.apply_preset(preset.id, run_id=run.id, page=1, page_size=200, selection=_selection())
    binding = _selection(displayed.selection_scope.watchlist_digest)
    query = _query_service(cache, _CapturedRankingStore(context, records, "ranking_active"))
    production, _ = query.export_projection(run.id, filters=MarketScanExportFilters(min_score=cutoff, selection=binding))
    spec = screen_spec_from_discovery(preset.criteria, preset.sort)
    filters = frozen_market_scan_export_filters(MarketScanExportFilters(selection=binding), spec.model_dump_json())
    frozen, research = query.export_projection(run.id, filters=filters)
    assert production.total != displayed.total
    assert frozen.total == displayed.total
    assert [item.symbol for item in frozen.items] == [item.symbol for item in displayed.items]
    assert [item.score for item in frozen.items] == [item.score for item in displayed.items]
    assert frozen.selection_scope == displayed.selection_scope
    assert frozen.production_ranking["status"] == "not_read"
    workbook = load_workbook(BytesIO(build_market_scan_workbook(frozen, filters,
        exported_at=datetime.fromisoformat("2026-10-16T12:00:00+08:00"), future_range_research=research).content))
    info = {row[0].value: row[1].value for row in workbook["导出信息"].iter_rows(min_row=2)}
    assert info["导出评分口径"] == "frozen"
    assert "未读取" in info["当前生产排名状态"]
    assert "当前生产评分规则" not in info and "概率生产排名状态" not in info
    assert json.loads(info["筛选合同 JSON"]) == spec.model_dump(mode="json")


@pytest.mark.parametrize("sort", ["market", "industry", "is_st", "is_new"])
def test_complete_saved_spec_export_supports_extra_sorts_and_two_sided_research_ranges(ranking_source, sort):
    cache, run, _context, records = ranking_source
    _members(cache, list(records))
    service = cache.domain_services.discovery
    preset = service.create_preset(DiscoveryPresetCreate(name=f"完整合同{sort}",
        criteria={"confidence": {"max": 100}, "risk": {"min": 0}, "tradability": {"max": 100}},
        sort=[{"field": sort, "order": "desc"}, {"field": "symbol", "order": "asc"}]))
    displayed = service.apply_preset(preset.id, run_id=run.id, page=1, page_size=200, selection=_selection())
    spec = screen_spec_from_discovery(preset.criteria, preset.sort)
    client = _client(cache, _query_service(cache))
    response = client.get(f"/api/market-scans/{run.id}/export.xlsx", params={"universe": "watchlist",
        "watchlist_digest": displayed.selection_scope.watchlist_digest, "ranking_basis": "frozen", "screen_spec": spec.model_dump_json()})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert [row[3].value for row in book["榜单"].iter_rows(min_row=2)] == [item.name for item in displayed.items]
    info = {row[0].value: row[1].value for row in book["导出信息"].iter_rows(min_row=2)}
    assert json.loads(info["筛选合同 JSON"]) == spec.model_dump(mode="json")


@pytest.mark.parametrize("params", [
    {"ranking_basis": "production", "screen_spec": "{}"},
    {"ranking_basis": "frozen"},
    {"ranking_basis": "frozen", "screen_spec": "{}", "sort": "rank"},
    {"ranking_basis": "frozen", "screen_spec": "{}", "min_confidence": "20"},
    {"ranking_basis": "frozen", "screen_spec": '{"unknown":1}'},
    {"ranking_basis": "frozen", "screen_spec": "not json"},
    {"ranking_basis": "frozen", "screen_spec": " " * 16385},
    {"ranking_basis": "unknown"},
])
def test_export_basis_and_complete_spec_are_strict_not_silently_ignored(ranking_source, params):
    cache, run, _context, _records = ranking_source
    assert _client(cache, _query_service(cache)).get(f"/api/market-scans/{run.id}/export.xlsx", params=params).status_code == 422
