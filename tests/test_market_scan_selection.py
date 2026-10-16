"""Current membership is a pinned read overlay, not a rewrite of frozen rows."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
import sqlite3
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.api.deps import get_market_scan_heavy_read_admission, get_market_scanner
from app.api.market_scan_read_admission import MarketScanHeavyReadAdmission
from app.api.routes import discovery, market_scan
from app.models.discovery import DiscoveryPresetCreate
from app.models.market_scan_screening import MarketScanScreenEvaluateRequest, ScreenSpecV2
from app.models.market_scan_selection import MarketScanSelectionRequest
from app.repositories import market_scan_verified_read
from app.repositories.market_scan_selection import MarketScanSelectionConflict, intersect_selection
from app.services.market_scan_export import MarketScanExportFilters, build_market_scan_workbook, frozen_market_scan_export_filters
from app.services.market_scan_query_service import MarketScanQueryService
from tests.test_market_scan_snapshot_integrity import _empty_research_stores, _result_query
from tests.test_strategy_execution import _environment


@pytest.fixture
def selection_env(tmp_path):
    cache, _service, _strategy, run_id = _environment(tmp_path)
    query = MarketScanQueryService(cache, _empty_research_stores())
    original = query.results(run_id, **_result_query())
    return cache, query, run_id, original


def _members(cache, symbols):
    with sqlite3.connect(cache.path) as conn:
        conn.execute("DELETE FROM watchlist")
        conn.executemany("INSERT INTO watchlist(symbol,code,market,name,created_at,updated_at) VALUES (?,?,?,?,?,?)",
            [(symbol, symbol[:6], symbol[-2:], "测试自选", "2026-08-11T00:00:00Z", "2026-08-11T00:00:00Z") for symbol in symbols])


def _selection(digest=None):
    return MarketScanSelectionRequest(universe="watchlist", watchlist_digest=digest)


def _read(query, run_id, *, digest=None, page=1, **filters):
    return query.results(run_id, **{**_result_query(), "page": page, "page_size": 1, **filters}, selection=_selection(digest))


def test_scope_filters_whole_snapshot_before_total_pagination_explanation_and_export(selection_env):
    cache, query, run_id, original = selection_env
    symbols = [item.symbol for item in original.items[-2:]]
    _members(cache, [*symbols, "000099.SZ"])
    first = _read(query, run_id)
    assert first.total == 2 and first.page_count == 2
    assert first.items[0].symbol == symbols[0]
    digest = first.selection_scope.watchlist_digest
    assert first.selection_scope.member_count == 3
    second = _read(query, run_id, digest=digest, page=2)
    assert [item.symbol for item in second.items] == [symbols[1]]
    evaluation = query.evaluate_screen(run_id, MarketScanScreenEvaluateRequest(
        universe="watchlist", watchlist_digest=digest, spec=ScreenSpecV2(status=None), page=2, page_size=1))
    assert evaluation.population_count == evaluation.matched_count == 2
    assert evaluation.selection_scope == first.selection_scope
    assert evaluation.matched.items == second.items
    full, _research = query.export_projection(run_id, filters=MarketScanExportFilters(status=None, selection=_selection(digest)))
    assert [item.symbol for item in full.items] == symbols and full.total == 2
    assert query.results(run_id, **_result_query()).run.snapshot_digest == original.run.snapshot_digest


def test_empty_watchlist_is_zero_for_all_read_paths(selection_env):
    cache, query, run_id, _original = selection_env
    _members(cache, [])
    first = _read(query, run_id)
    assert (first.total, first.page_count, first.items, first.selection_scope.member_count) == (0, 0, [], 0)
    digest = first.selection_scope.watchlist_digest
    evaluated = query.evaluate_screen(run_id, MarketScanScreenEvaluateRequest(universe="watchlist", watchlist_digest=digest))
    assert evaluated.population_count == evaluated.matched_count == 0
    assert evaluated.matched.items == evaluated.near_misses == []
    exported, _ = query.export_projection(run_id, filters=MarketScanExportFilters(selection=_selection(digest)))
    assert exported.total == 0 and exported.items == []


def test_large_historical_membership_does_not_truncate_current_batch_member(selection_env):
    cache, query, run_id, original = selection_env
    symbol = max(item.symbol for item in original.items)
    members = [f"{index:06}.SZ" for index in range(10001)] + [symbol]
    assert sorted(members).index(symbol) == 10001
    _members(cache, members)
    page = _read(query, run_id)
    assert page.total == 1 and [item.symbol for item in page.items] == [symbol]
    assert page.selection_scope.member_count == 10002
    digest = page.selection_scope.watchlist_digest
    spec = ScreenSpecV2(status=None)
    evaluation = query.evaluate_screen(run_id, MarketScanScreenEvaluateRequest(
        universe="watchlist", watchlist_digest=digest, spec=spec, page=1, page_size=1))
    assert evaluation.population_count == evaluation.matched_count == 1
    assert evaluation.matched.items == page.items
    filters = MarketScanExportFilters(status=None, selection=_selection(digest))
    for export_filter in (filters, frozen_market_scan_export_filters(filters, spec.model_dump_json())):
        exported, _ = query.export_projection(run_id, filters=export_filter)
        assert exported.total == 1 and [item.symbol for item in exported.items] == [symbol]
        assert exported.selection_scope == page.selection_scope


def test_membership_changes_reject_every_pinned_path_until_explicit_refresh(selection_env):
    cache, query, run_id, original = selection_env
    _members(cache, [item.symbol for item in original.items[:2]])
    old = _read(query, run_id)
    digest = old.selection_scope.watchlist_digest
    _members(cache, [original.items[-1].symbol])
    operations = [lambda: _read(query, run_id, digest=digest, page=2),
        lambda: query.evaluate_screen(run_id, MarketScanScreenEvaluateRequest(universe="watchlist", watchlist_digest=digest)),
        lambda: query.export_projection(run_id, filters=MarketScanExportFilters(selection=_selection(digest)))]
    for operation in operations:
        with pytest.raises(MarketScanSelectionConflict, match="刷新自选范围"):
            operation()
    new = _read(query, run_id)
    assert new.selection_scope.watchlist_digest != digest
    assert new.total == 1 and new.items[0].symbol == original.items[-1].symbol


def test_unpinned_continuations_are_rejected_and_nonmember_filters_intersect(selection_env):
    cache, query, run_id, original = selection_env
    _members(cache, [original.items[-1].symbol])
    assert _read(query, run_id, keyword=original.items[0].symbol[:6]).total == 0
    for operation in [lambda: _read(query, run_id, page=2),
        lambda: query.evaluate_screen(run_id, MarketScanScreenEvaluateRequest(universe="watchlist")),
        lambda: query.export_projection(run_id, filters=MarketScanExportFilters(selection=_selection()))]:
        with pytest.raises(MarketScanSelectionConflict):
            operation()
    assert intersect_selection(("a", "b"), ("b", "c")) == ("b",)
    assert intersect_selection(("a",), ()) == ()


def test_membership_and_results_share_one_read_snapshot(selection_env, monkeypatch):
    cache, query, run_id, original = selection_env
    first_symbol, second_symbol = [item.symbol for item in original.items[:2]]
    _members(cache, [first_symbol])
    actual = market_scan_verified_read.watchlist_symbols
    def mutate_after_membership_read(conn):
        symbols = actual(conn)
        _members(cache, [second_symbol])
        return symbols
    monkeypatch.setattr(market_scan_verified_read, "watchlist_symbols", mutate_after_membership_read)
    page = _read(query, run_id)
    assert [item.symbol for item in page.items] == [first_symbol]
    monkeypatch.setattr(market_scan_verified_read, "watchlist_symbols", actual)
    with pytest.raises(MarketScanSelectionConflict):
        _read(query, run_id, digest=page.selection_scope.watchlist_digest)


def test_saved_preset_overlay_preserves_definition_and_member_pin(selection_env):
    cache, query, run_id, original = selection_env
    service = cache.domain_services.discovery
    preset = service.create_preset(DiscoveryPresetCreate(name="全部有效", criteria={}, sort=[{"field": "rank", "order": "asc"}]))
    _members(cache, [item.symbol for item in original.items[-2:]])
    first = service.apply_preset(preset.id, run_id=run_id, page=1, page_size=1, selection=_selection())
    digest = first.selection_scope.watchlist_digest
    second = service.apply_preset(preset.id, run_id=run_id, page=2, page_size=1, selection=_selection(digest))
    assert first.total == second.total == 2 and second.items[0].position == 2
    assert second.items[0].symbol == _read(query, run_id, page=2, digest=digest).items[0].symbol
    assert service.get_preset(preset.id) == preset
    _members(cache, [])
    with pytest.raises(MarketScanSelectionConflict):
        service.apply_preset(preset.id, run_id=run_id, page=2, page_size=1, selection=_selection(digest))
    assert service.apply_preset(preset.id, run_id=run_id, page=1, page_size=1, selection=_selection()).total == 0


def _client(cache, query):
    app = FastAPI()
    app.include_router(market_scan.router)
    app.include_router(discovery.router)
    def export_results(run_id, filters):
        page, research = query.export_projection(run_id, filters=filters)
        return build_market_scan_workbook(page, filters, exported_at=datetime.fromisoformat("2026-10-16T12:00:00+08:00"), future_range_research=research)
    scanner = SimpleNamespace(results=query.results, evaluate_screen=query.evaluate_screen, export_results=export_results)
    app.dependency_overrides[get_market_scanner] = lambda: scanner
    app.dependency_overrides[get_market_scan_heavy_read_admission] = lambda: MarketScanHeavyReadAdmission()
    app.dependency_overrides[discovery.get_discovery_service] = lambda: cache.domain_services.discovery
    return TestClient(app)


def test_http_scope_routes_conflict_and_workbook_metadata(selection_env):
    cache, query, run_id, original = selection_env
    _members(cache, [original.items[-1].symbol])
    client = _client(cache, query)
    response = client.get(f"/api/market-scans/{run_id}/results?universe=watchlist")
    assert response.status_code == 200 and response.json()["total"] == 1
    digest = response.json()["selection_scope"]["watchlist_digest"]
    exported = client.get(f"/api/market-scans/{run_id}/export.xlsx?universe=watchlist&watchlist_digest={digest}")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content))
    info = {row[0].value: row[1].value for row in workbook["导出信息"].iter_rows(min_row=2)}
    assert info["自选成员摘要"] == digest and info["当前查看范围"] == "当前自选"
    assert workbook["榜单"].max_row == 2
    _members(cache, [])
    for path in ("results", "export.xlsx"):
        stale = client.get(f"/api/market-scans/{run_id}/{path}?universe=watchlist&watchlist_digest={digest}")
        assert stale.status_code == 409 and stale.headers["cache-control"] == "no-store"
    stale = client.post(f"/api/market-scans/{run_id}/screen/evaluate", json={"universe": "watchlist", "watchlist_digest": digest})
    assert stale.status_code == 409 and stale.headers["cache-control"] == "no-store"


@pytest.mark.parametrize("params", ["universe=oops", "universe=all&watchlist_digest=" + "a" * 64, "universe=watchlist&watchlist_digest=short"])
def test_http_selection_rejects_invalid_binding(selection_env, params):
    cache, query, run_id, _original = selection_env
    assert _client(cache, query).get(f"/api/market-scans/{run_id}/results?{params}").status_code == 422
